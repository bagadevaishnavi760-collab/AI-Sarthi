import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from pydantic import ValidationError
from api.ollama_client import OllamaUnavailable
from api.rag_service import (
    GeneratedAnswer,
    IndexUnavailable,
    RAGError,
    RAGService,
    _extract_document,
)


class FakeStructuredData:
    def __init__(self, result=None):
        self.result = result or {"records": [], "data_sources": [], "calculations": [], "limitations": []}

    def query(self, question):
        del question
        return self.result


class FakeIndex:
    def __init__(self, matches):
        self.matches = matches
        self.search_options = {}

    def similarity_search_with_score(self, question, k, **kwargs):
        del question, k
        self.search_options = kwargs
        return self.matches


class FakeDocument:
    def __init__(self, page_content, metadata):
        self.page_content = page_content
        self.metadata = metadata


class FakeSplitter:
    def __init__(self, **kwargs):
        del kwargs

    def split_documents(self, documents):
        return documents


class FakeFAISS:
    fail_creation = False
    last_documents = []

    @classmethod
    def from_documents(cls, documents, embeddings):
        del embeddings
        if cls.fail_creation:
            raise RuntimeError("simulated FAISS write failure")
        cls.last_documents = documents
        return cls()

    @classmethod
    def load_local(cls, folder, embeddings, allow_dangerous_deserialization):
        del embeddings, allow_dangerous_deserialization
        if not (Path(folder) / "index.faiss").is_file():
            raise OSError("missing test index")
        return cls()

    def save_local(self, folder):
        destination = Path(folder)
        destination.mkdir(parents=True, exist_ok=True)
        (destination / "index.faiss").write_bytes(b"test index")
        (destination / "index.pkl").write_bytes(b"test metadata")


class RAGServiceTests(unittest.TestCase):
    def test_retrieved_evidence_is_in_generation_context_and_citation_is_real(self):
        service = RAGService(FakeStructuredData())
        excerpt = "The programme report states the verified release amount."
        document = SimpleNamespace(
            page_content=excerpt,
            metadata={
                "title": "NHM report",
                "filename": "nhm.pdf",
                "page": 4,
                "section": "Financial progress",
                "source_url": None,
            },
        )
        service._load_index = lambda: (FakeIndex([(document, 0.25)]), {})
        captured = {}

        def generate(prompt):
            captured["prompt"] = prompt
            return "The report records the release amount."

        service._generate = generate
        result = service.answer("What does the report say about release?")

        self.assertTrue(result["grounded"])
        self.assertIn(excerpt, captured["prompt"])
        self.assertEqual(result["sources"][0]["page"], 4)
        self.assertEqual(result["sources"][0]["excerpt"], excerpt)
        self.assertEqual(result["answer"], "The report records the release amount.")

    def test_no_relevant_evidence_returns_insufficient_answer_without_generation(self):
        service = RAGService(FakeStructuredData())
        service._load_index = lambda: (FakeIndex([]), {})
        service._generate = lambda prompt: self.fail("Generation must not run without evidence.")

        result = service.answer("Question with no matching source")

        self.assertFalse(result["grounded"])
        self.assertEqual(result["sources"], [])
        self.assertIn("do not establish", result["answer"])

    def test_distance_threshold_discards_irrelevant_retrieval(self):
        service = RAGService(FakeStructuredData())
        document = SimpleNamespace(page_content="unrelated", metadata={"title": "doc"})
        service._load_index = lambda: (FakeIndex([(document, 2.0)]), {})
        service._generate = lambda prompt: self.fail("Irrelevant evidence must be discarded.")

        result = service.answer("A query unrelated to this chunk")

        self.assertFalse(result["grounded"])

    def test_retrieval_supports_metadata_filters(self):
        service = RAGService(FakeStructuredData())
        document = SimpleNamespace(
            page_content="filtered evidence",
            metadata={"title": "doc", "filename": "doc.pdf", "document_type": "pdf"},
        )
        index = FakeIndex([(document, 0.2)])
        service._load_index = lambda: (index, {})
        service._generate = lambda prompt: "Grounded answer."

        with patch.dict("os.environ", {"RAG_RETRIEVAL_K": "5"}):
            service.answer("Find this document", metadata_filter={"document_type": "pdf"})

        self.assertEqual(index.search_options["filter"], {"document_type": "pdf"})
        self.assertEqual(index.search_options["fetch_k"], 20)

    def test_missing_index_does_not_block_available_structured_evidence(self):
        structured = FakeStructuredData({
            "records": [{
                "State": "Bihar",
                "Year": "2022-23",
                "Expenditure": 75,
                "provenance": "budget.csv",
                "state": "Bihar",
                "year": "2022-23",
            }],
            "data_sources": [{
                "filename": "budget.csv",
                "dataset": "NHM budget",
                "provenance": "Project CSV: budget.csv",
                "metrics": [],
            }],
            "calculations": [],
            "limitations": [],
        })
        service = RAGService(structured)
        service._load_index = lambda: (_ for _ in ()).throw(IndexUnavailable("index missing"))
        service._generate = lambda prompt: "The structured record reports expenditure of 75."

        result = service.answer("NHM expenditure in Bihar")

        self.assertTrue(result["grounded"])
        self.assertEqual(result["data_sources"][0]["states"], ["Bihar"])
        self.assertEqual(result["data_sources"][0]["years"], ["2022-23"])
        self.assertIn("index missing", result["limitations"])

    def test_pdf_page_extraction_preserves_page_numbers(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "report.pdf"
            writer = PdfWriter()
            for text in ("Page one evidence", "Page two evidence"):
                page = writer.add_blank_page(width=300, height=300)
                font = DictionaryObject({
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                })
                page[NameObject("/Resources")] = DictionaryObject({
                    NameObject("/Font"): DictionaryObject({
                        NameObject("/F1"): writer._add_object(font),
                    }),
                })
                stream = DecodedStreamObject()
                stream.set_data(f"BT /F1 12 Tf 72 200 Td ({text}) Tj ET".encode())
                page[NameObject("/Contents")] = writer._add_object(stream)
            with path.open("wb") as output:
                writer.write(output)
            pages = _extract_document(path)

        self.assertEqual(pages, [(1, "Page one evidence"), (2, "Page two evidence")])

    def test_docx_table_text_is_preserved(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "table.docx"
            document = Document()
            document.add_heading("Facility summary", level=1)
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "PHCs"
            table.rows[0].cells[1].text = "25"
            document.save(path)

            pages = _extract_document(path)

        self.assertIn("Facility summary", pages[0][1])
        self.assertIn("PHCs | 25", pages[0][1])

    def test_empty_and_unsupported_documents_are_not_indexable(self):
        with tempfile.TemporaryDirectory() as folder:
            empty = Path(folder) / "empty.txt"
            empty.write_text("", encoding="utf-8")
            unsupported = Path(folder) / "source.csv"
            unsupported.write_text("x", encoding="utf-8")

            self.assertEqual(_extract_document(empty), [(None, "")])
            with self.assertRaisesRegex(RAGError, "Unsupported document format"):
                _extract_document(unsupported)

    def test_unreadable_text_document_reports_error(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "unreadable.txt"
            path.write_bytes(b"\xff\xfe\xff")
            with self.assertRaisesRegex(RAGError, "Could not extract text"):
                _extract_document(path)

    def test_index_persistence_idempotency_and_failed_update_safety(self):
        documents_temp = tempfile.TemporaryDirectory()
        index_temp = tempfile.TemporaryDirectory()
        self.addCleanup(documents_temp.cleanup)
        self.addCleanup(index_temp.cleanup)
        documents_dir = Path(documents_temp.name)
        index_dir = Path(index_temp.name)
        (documents_dir / "policy.md").write_text("1. POLICY\nOriginal evidence.", encoding="utf-8")

        core_documents = types.ModuleType("langchain_core.documents")
        core_documents.Document = FakeDocument
        splitters = types.ModuleType("langchain_text_splitters")
        splitters.RecursiveCharacterTextSplitter = FakeSplitter
        community = types.ModuleType("langchain_community")
        community.__path__ = []
        vectorstores = types.ModuleType("langchain_community.vectorstores")
        vectorstores.FAISS = FakeFAISS
        modules = {
            "langchain_core.documents": core_documents,
            "langchain_text_splitters": splitters,
            "langchain_community": community,
            "langchain_community.vectorstores": vectorstores,
        }
        service = RAGService(FakeStructuredData())
        service.documents_dir = documents_dir
        service.index_dir = index_dir
        service._get_embeddings = lambda: object()
        FakeFAISS.fail_creation = False
        self.addCleanup(setattr, FakeFAISS, "fail_creation", False)
        with patch.dict(sys.modules, modules):
            first = service.ingest()
            first_pointer = json.loads((index_dir / "current.json").read_text(encoding="utf-8"))
            self.assertEqual(first["status"], "indexed")
            self.assertEqual(service._current_manifest()[1]["chunk_count"], 1)
            metadata = FakeFAISS.last_documents[0].metadata
            indexed_documents = service._current_manifest()[1]["documents"]
            self.assertIn(metadata["document_id"], indexed_documents)
            self.assertEqual(
                metadata["document_id"],
                indexed_documents[metadata["document_id"]]["document_id"],
            )
            self.assertEqual(metadata["document_type"], "md")
            self.assertEqual(metadata["section"], "1. POLICY")
            self.assertTrue(metadata["chunk_id"])
            self.assertEqual(service.ingest()["status"], "unchanged")

            (documents_dir / "policy.md").write_text("1. POLICY\nModified evidence.", encoding="utf-8")
            FakeFAISS.fail_creation = True
            with self.assertRaisesRegex(RAGError, "previous index was preserved"):
                service.ingest()
            self.assertEqual(
                json.loads((index_dir / "current.json").read_text(encoding="utf-8")),
                first_pointer,
            )

            FakeFAISS.fail_creation = False
            updated = service.ingest()
            second_pointer = json.loads((index_dir / "current.json").read_text(encoding="utf-8"))
            self.assertEqual(updated["status"], "indexed")
            self.assertNotEqual(first_pointer, second_pointer)
            self.assertTrue((index_dir / first_pointer["generation"]).is_dir())
            self.assertTrue(service._load_index())

    def test_missing_and_incompatible_indices_are_reported(self):
        with tempfile.TemporaryDirectory() as folder:
            service = RAGService(FakeStructuredData())
            service.index_dir = Path(folder)
            with self.assertRaisesRegex(IndexUnavailable, "No FAISS index"):
                service._current_manifest()

            generation_name = "generation-test"
            generation = service.index_dir / generation_name
            generation.mkdir()
            (generation / "index.faiss").write_bytes(b"index")
            (generation / "index.pkl").write_bytes(b"metadata")
            (generation / "manifest.json").write_text(
                json.dumps({
                    "embedding_model": "different-model",
                    "chunk_count": 0,
                    "documents": {},
                    "indexed_at": "test",
                }),
                encoding="utf-8",
            )
            (service.index_dir / "current.json").write_text(
                json.dumps({"generation": generation_name}), encoding="utf-8"
            )
            with self.assertRaisesRegex(IndexUnavailable, "different embedding model"):
                service._current_manifest()

    def test_health_does_not_equate_api_availability_with_rag_readiness(self):
        service = RAGService(FakeStructuredData())
        service._embedding_model_cached = lambda: False
        service._current_manifest = lambda: (_ for _ in ()).throw(IndexUnavailable("No index"))
        service.ollama.tags = lambda: (_ for _ in ()).throw(OllamaUnavailable("offline"))
        status = service.health()

        self.assertFalse(status["ready"])
        self.assertFalse(status["embedding_model_available"])
        self.assertFalse(status["index_available"])
        self.assertFalse(status["ollama_connected"])
        self.assertIn("faiss", status["missing_rag_dependencies"])

    def test_model_output_has_one_bounded_retry_and_validates_json(self):
        class FakeOllama:
            def __init__(self, **kwargs):
                self.calls = 0

            def chat(self, messages, **kwargs):
                del messages, kwargs
                self.calls += 1
                return "not json" if self.calls == 1 else '{"answer":" grounded answer "}'

        service = RAGService()
        service.ollama = FakeOllama()
        answer = service._generate("prompt")

        self.assertEqual(answer, "grounded answer")
        with self.assertRaises(ValidationError):
            GeneratedAnswer.model_validate_json('{"answer":" "}')

    def test_ollama_inference_failure_is_explicit(self):
        service = RAGService()
        service.ollama.chat = lambda *args, **kwargs: (_ for _ in ()).throw(
            OllamaUnavailable("offline")
        )
        with self.assertRaises(OllamaUnavailable):
            service._generate("prompt")


if __name__ == "__main__":
    unittest.main()
